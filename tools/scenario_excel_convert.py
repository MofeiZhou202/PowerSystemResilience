#!/usr/bin/env python3
"""Convert generated scenario JSON bundles to/from editable Excel workbooks.

The workbook is intentionally scenario-oriented rather than a full system model
spreadsheet: SYSTEM_JSON/RAW_JSON preserve lossless round-trip data, while
PROFILES, BINDINGS, NORMALIZATION, and DEVICE_CAPACITY expose the fields users
need to inspect or edit.
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
from pathlib import Path
from typing import Any

try:
    import openpyxl
except Exception as exc:  # pragma: no cover - surfaced in server error
    raise SystemExit(f"openpyxl is required for scenario Excel conversion: {exc}")

WORKBOOK_SCHEMA = "HACDCPF_SCENARIO_WORKBOOK_V1"


def json_dumps(obj: Any) -> str:
    return json.dumps(obj, ensure_ascii=False, separators=(",", ":"))


def normalize_case(case: dict[str, Any], index: int) -> dict[str, Any]:
    if "system" in case or "generated_scenario" in case or "standard_time_series" in case:
        system = dict(case.get("system") or {})
        meta = dict(case.get("generated_scenario") or case.get("_generated_scenario") or {})
        ts = dict(case.get("standard_time_series") or case.get("_time_series") or {})
        case_id = case.get("case_id") or meta.get("representative_id") or meta.get("scenario_id") or f"case_{index + 1}"
        return {"case_id": case_id, "system": system, "generated_scenario": meta, "standard_time_series": ts, "raw_case": case}

    system = dict(case)
    meta = dict(system.pop("_generated_scenario", {}) or {})
    ts = dict(system.pop("_time_series", {}) or {})
    case_id = meta.get("representative_id") or meta.get("scenario_id") or system.get("name") or f"case_{index + 1}"
    return {"case_id": case_id, "system": system, "generated_scenario": meta, "standard_time_series": ts, "raw_case": case}


def normalize_bundle(obj: dict[str, Any]) -> dict[str, Any]:
    if "cases" not in obj and ("ac" in obj or "dc" in obj or "system" in obj):
        obj = {"format": "generated_scenario_case_bundle_v2", "schema_version": 2, "cases": [obj]}
    cases = [normalize_case(c, i) for i, c in enumerate(obj.get("cases") or [])]
    return {
        "format": obj.get("format", "generated_scenario_case_bundle_v2"),
        "schema_version": obj.get("schema_version", 2),
        "unit_space": obj.get("unit_space", "dimensionless_multiplier"),
        "family": obj.get("family") or (cases[0]["generated_scenario"].get("family") if cases else "unknown"),
        "case_count": len(cases),
        "source": obj,
        "cases": cases,
    }


def append_row(ws, values):
    ws.append(list(values))


def profiles_for(case: dict[str, Any]) -> list[dict[str, Any]]:
    # Only expose profiles that have direct runtime meaning. The aggregate
    # renewable profile duplicates PV+wind and confused users, so keep PV and
    # wind separate and omit scenario_renewable_scale from editable workbooks.
    out = []
    for p in list((case.get("standard_time_series") or {}).get("profiles") or []):
        if str(p.get("name") or "") == "scenario_renewable_scale":
            continue
        out.append(p)
    return out


def profile_col_name(profile: dict[str, Any], fallback: int) -> str:
    pid = profile.get("id", fallback)
    name = str(profile.get("name") or "profile")
    safe = re.sub(r"[^0-9A-Za-z_\-一-鿿]+", "_", name)[:80]
    return f"p{pid}_{safe}"


def profile_id_from_header(header: str) -> int | None:
    m = re.match(r"p(-?\d+)_", str(header or ""))
    return int(m.group(1)) if m else None


def service_on(item: dict[str, Any]) -> bool:
    return item.get("in_service", True) is not False


def first_positive(item: dict[str, Any], keys: list[str]) -> tuple[float, str]:
    for key in keys:
        try:
            val = float(item.get(key) or 0.0)
        except Exception:
            val = 0.0
        if val > 1e-9:
            return val, key
    return 0.0, keys[0]


def text_has(item: dict[str, Any], token: str) -> bool:
    text = f"{item.get('type','')} {item.get('sgen_type','')} {item.get('name','')}".lower()
    return token.lower() in text


def iter_device_rows(case_id: str, system: dict[str, Any]):
    ac = system.get("ac") or {}
    dc = system.get("dc") or {}

    def emit(kind: str, items: list[dict[str, Any]], keys: list[str], energy_key: str | None = None):
        for pos, item in enumerate(items or []):
            if not service_on(item):
                continue
            base_mw, field = first_positive(item, keys)
            base_mwh = float(item.get(energy_key, 0.0) or 0.0) if energy_key else 0.0
            if base_mw <= 1e-9 and base_mwh <= 1e-9:
                continue
            yield [case_id, kind, item.get("index", pos + 1), pos, item.get("bus", 0), item.get("name", ""), base_mw, item.get("q_mvar", 0.0), base_mwh, field, service_on(item), item.get("profile_id", -1), True]

    ac_loads = ac.get("loads") or []
    dc_loads = dc.get("loads") or []
    if ac_loads:
        yield from emit("AC_LOAD", ac_loads, ["p_mw"])
    else:
        # MATPOWER / dist33-style cases often store demand on buses instead of
        # explicit load components. Expose those bus demands as editable capacity
        # rows so changing base_mw actually changes the imported case.
        yield from emit("AC_BUS", ac.get("buses") or [], ["pd_mw"])
    if dc_loads:
        yield from emit("DC_LOAD", dc_loads, ["p_mw", "p_rated_mw"])
    else:
        yield from emit("DC_BUS", dc.get("buses") or [], ["pd_mw"])
    yield from emit("AC_PV_SYSTEM", ac.get("pv_systems") or [], ["pmax_mw", "p_mw", "sn_mva"])
    yield from emit("DC_PV_ARRAY", dc.get("pv_arrays") or [], ["p_set_mw"])
    for pos, item in enumerate(ac.get("renewable_gens") or []):
        if not service_on(item):
            continue
        base, field = first_positive(item, ["p_rated_mw", "p_mw"])
        if base > 1e-9:
            yield [case_id, "AC_WIND" if text_has(item, "wind") else "AC_RENEWABLE", item.get("index", pos + 1), pos, item.get("bus", 0), item.get("name", ""), base, 0.0, 0.0, field, True, item.get("profile_id", -1), True]
    for pos, item in enumerate(ac.get("static_generators") or []):
        if service_on(item) and (text_has(item, "pv") or text_has(item, "solar") or text_has(item, "wind")):
            base, field = first_positive(item, ["p_rated_mw", "pmax_mw", "p_mw"])
            if base > 1e-9:
                yield [case_id, "AC_STATIC_WIND" if text_has(item, "wind") else "AC_STATIC_PV", item.get("index", pos + 1), pos, item.get("bus", 0), item.get("name", ""), base, 0.0, 0.0, field, True, item.get("profile_id", -1), True]
    for pos, item in enumerate(dc.get("dc_static_generators") or []):
        if service_on(item) and (text_has(item, "pv") or text_has(item, "solar") or text_has(item, "wind")):
            base, field = first_positive(item, ["p_set_mw"])
            if base > 1e-9:
                yield [case_id, "DC_STATIC_WIND" if text_has(item, "wind") else "DC_STATIC_PV", item.get("index", pos + 1), pos, item.get("bus", 0), item.get("name", ""), base, 0.0, 0.0, field, True, item.get("profile_id", -1), True]
    yield from emit("AC_STORAGE", ac.get("storage") or [], ["p_rated_mw", "pmax_mw", "p_mw"], "e_rated_mwh")


def json_to_xlsx(json_path: Path, xlsx_path: Path) -> None:
    bundle = normalize_bundle(json.loads(json_path.read_text(encoding="utf-8")))
    wb = openpyxl.Workbook()
    wb.active.title = "METADATA"
    ws = wb["METADATA"]
    append_row(ws, ["key", "value"])
    for k, v in [
        ("workbook_schema", WORKBOOK_SCHEMA),
        ("json_format", bundle["format"]),
        ("schema_version", bundle["schema_version"]),
        ("unit_space", bundle["unit_space"]),
        ("family", bundle["family"]),
        ("case_count", bundle["case_count"]),
    ]:
        append_row(ws, [k, v])

    ws = wb.create_sheet("SCENARIOS")
    append_row(ws, ["case_id", "scenario_id", "family", "cluster_id", "probability", "member_count", "representative_id", "intensity", "notes"])
    for case in bundle["cases"]:
        meta = case["generated_scenario"]
        append_row(ws, [case["case_id"], meta.get("scenario_id", ""), meta.get("family", bundle["family"]), meta.get("cluster_id", 0), meta.get("probability", 0), meta.get("member_count", 0), meta.get("representative_id", ""), meta.get("intensity", ""), ""])

    ws = wb.create_sheet("PROFILE_CATALOG")
    append_row(ws, ["case_id", "profile_id", "profile_name", "domain", "scope", "unit", "description"])
    for case in bundle["cases"]:
        for p in profiles_for(case):
            append_row(ws, [case["case_id"], p.get("id"), p.get("name", ""), p.get("domain", ""), p.get("scope", ""), p.get("unit", "multiplier"), p.get("description", "")])

    ws = wb.create_sheet("PROFILES")
    first_profiles = profiles_for(bundle["cases"][0]) if bundle["cases"] else []
    append_row(ws, ["case_id", "time_index", "hour", *[profile_col_name(p, i) for i, p in enumerate(first_profiles)]])
    for case in bundle["cases"]:
        ts = case.get("standard_time_series") or {}
        profs = profiles_for(case)
        num_steps = int(ts.get("num_steps") or (len(profs[0].get("values") or []) if profs else 0))
        step = float(ts.get("step_duration_hr") or 1.0)
        for t in range(num_steps):
            row = [case["case_id"], t, t * step]
            for p in profs:
                vals = p.get("values") or []
                row.append(vals[t] if t < len(vals) else None)
            append_row(ws, row)

    ws = wb.create_sheet("BINDINGS")
    append_row(ws, ["case_id", "binding_type", "kind", "component_index", "position", "bus", "profile_id", "target_field", "base_mw", "note"])
    for case in bundle["cases"]:
        binding = (case.get("standard_time_series") or {}).get("binding") or {}
        append_row(ws, [case["case_id"], "assign_all_loads", "", "", "", "", binding.get("assign_all_loads_to", -1), "", "", ""])
        append_row(ws, [case["case_id"], "assign_all_pv", "", "", "", "", binding.get("assign_all_pv_to", -1), "", "", ""])
        for row in binding.get("load_profile_map") or []:
            append_row(ws, [case["case_id"], "load_profile_map", row.get("kind", ""), row.get("load_index", ""), row.get("load_position", ""), row.get("bus", ""), row.get("profile_id", -1), "", row.get("base_mw", ""), ""])

    ws = wb.create_sheet("NORMALIZATION")
    append_row(ws, ["case_id", "base_load_mw", "base_pv_mw", "base_wind_mw", "base_other_renewable_mw", "base_renewable_mw", "base_storage_mwh", "num_steps", "step_duration_hr", "profile_semantics"])
    for case in bundle["cases"]:
        ts = case.get("standard_time_series") or {}
        n = ts.get("normalization") or {}
        append_row(ws, [case["case_id"], n.get("base_load_mw", 0), n.get("base_pv_mw", 0), n.get("base_wind_mw", 0), n.get("base_other_renewable_mw", 0), n.get("base_renewable_mw", 0), n.get("base_storage_mwh", 0), ts.get("num_steps", 0), ts.get("step_duration_hr", 1), n.get("profile_semantics", "dimensionless_multiplier")])

    ws = wb.create_sheet("DEVICE_CAPACITY")
    append_row(ws, ["case_id", "kind", "component_index", "position", "bus", "name", "base_mw", "base_mvar", "base_mwh", "capacity_field", "in_service", "profile_id", "editable"])
    for case in bundle["cases"]:
        for row in iter_device_rows(case["case_id"], case["system"]):
            append_row(ws, row)

    ws = wb.create_sheet("SYSTEM_JSON")
    append_row(ws, ["case_id", "chunk_index", "json_chunk"])
    for case in bundle["cases"]:
        append_row(ws, [case["case_id"], 0, json_dumps(case["system"])])

    ws = wb.create_sheet("RAW_JSON")
    append_row(ws, ["chunk_index", "json_chunk"])
    raw = json_dumps(bundle["source"])
    chunk = 30000
    for i in range(0, len(raw), chunk):
        append_row(ws, [i // chunk, raw[i : i + chunk]])

    ws = wb.create_sheet("WARNINGS")
    append_row(ws, ["case_id", "severity", "message"])
    wb.save(xlsx_path)


def sheet_rows(wb, name: str):
    if name not in wb.sheetnames:
        return []
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h or "") for h in rows[0]]
    out = []
    for row in rows[1:]:
        if not any(v is not None and v != "" for v in row):
            continue
        out.append({headers[i]: row[i] if i < len(row) else None for i in range(len(headers))})
    return out


def find_device(system: dict[str, Any], kind: str, index: int, position: int):
    ac = system.setdefault("ac", {})
    dc = system.setdefault("dc", {})
    arrays = {
        "AC_LOAD": ac.setdefault("loads", []),
        "AC_BUS": ac.setdefault("buses", []),
        "DC_LOAD": dc.setdefault("loads", []),
        "DC_BUS": dc.setdefault("buses", []),
        "AC_PV_SYSTEM": ac.setdefault("pv_systems", []),
        "DC_PV_ARRAY": dc.setdefault("pv_arrays", []),
        "AC_WIND": ac.setdefault("renewable_gens", []),
        "AC_RENEWABLE": ac.setdefault("renewable_gens", []),
        "AC_STATIC_WIND": ac.setdefault("static_generators", []),
        "AC_STATIC_PV": ac.setdefault("static_generators", []),
        "DC_STATIC_WIND": dc.setdefault("dc_static_generators", []),
        "DC_STATIC_PV": dc.setdefault("dc_static_generators", []),
        "AC_STORAGE": ac.setdefault("storage", []),
    }
    arr = arrays.get(kind)
    if not isinstance(arr, list):
        return None
    for item in arr:
        if int(item.get("index", -999999)) == index:
            return item
    if 0 <= position < len(arr):
        return arr[position]
    return None


def xlsx_to_json(xlsx_path: Path, json_path: Path) -> None:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    raw_rows = sheet_rows(wb, "RAW_JSON")
    if raw_rows:
        raw = "".join(str(r.get("json_chunk") or "") for r in sorted(raw_rows, key=lambda r: int(r.get("chunk_index") or 0)))
        bundle = json.loads(raw)
    else:
        system_by_case = {}
        for r in sheet_rows(wb, "SYSTEM_JSON"):
            system_by_case[str(r.get("case_id"))] = json.loads(r.get("json_chunk") or "{}")
        cases = []
        for r in sheet_rows(wb, "SCENARIOS"):
            cid = str(r.get("case_id"))
            system = system_by_case.get(cid, {})
            meta = {"family": r.get("family"), "scenario_id": r.get("scenario_id"), "representative_id": r.get("representative_id"), "cluster_id": r.get("cluster_id"), "probability": r.get("probability"), "member_count": r.get("member_count")}
            case = dict(system)
            case["_generated_scenario"] = {k: v for k, v in meta.items() if v not in (None, "")}
            case["_time_series"] = {"profiles": []}
            cases.append(case)
        bundle = {"format": "generated_scenario_case_bundle_v2", "schema_version": 2, "cases": cases, "case_count": len(cases)}

    # Full /api/session/generate_scenarios exports are raw result JSON, not
    # importable generated-case bundles. Preserve those losslessly via RAW_JSON.
    if "cases" not in bundle and not any(k in bundle for k in ("ac", "dc", "system")):
        json_path.write_text(json.dumps(bundle, ensure_ascii=False, indent=2), encoding="utf-8")
        return

    norm = normalize_bundle(bundle)
    case_by_id = {c["case_id"]: c for c in norm["cases"]}

    catalog = {}
    for r in sheet_rows(wb, "PROFILE_CATALOG"):
        cid = str(r.get("case_id"))
        pid = int(r.get("profile_id") or 0)
        catalog.setdefault(cid, {})[pid] = {"id": pid, "name": r.get("profile_name") or "", "domain": r.get("domain") or "", "scope": r.get("scope") or "", "unit": r.get("unit") or "multiplier", "values": []}

    if "PROFILES" in wb.sheetnames:
        ws = wb["PROFILES"]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        id_by_col = {i: profile_id_from_header(h) for i, h in enumerate(headers)}
        for row in ws.iter_rows(min_row=2, values_only=True):
            cid = str(row[0] or "")
            if not cid:
                continue
            for i, pid in id_by_col.items():
                if pid is None:
                    continue
                catalog.setdefault(cid, {}).setdefault(pid, {"id": pid, "name": str(headers[i]), "unit": "multiplier", "values": []})["values"].append(float(row[i] or 0.0))

    binding_by_case = {}
    for r in sheet_rows(wb, "BINDINGS"):
        cid = str(r.get("case_id"))
        b = binding_by_case.setdefault(cid, {"load_profile_map": []})
        typ = r.get("binding_type")
        pid = int(r.get("profile_id") or -1)
        if typ == "assign_all_loads":
            b["assign_all_loads_to"] = pid
        elif typ == "assign_all_pv":
            b["assign_all_pv_to"] = pid
        elif typ == "load_profile_map":
            row = {"kind": r.get("kind"), "profile_id": pid, "bus": r.get("bus")}
            if r.get("component_index") not in (None, ""):
                row["load_index"] = int(r.get("component_index"))
            if r.get("position") not in (None, ""):
                row["load_position"] = int(r.get("position"))
            b["load_profile_map"].append(row)

    norm_by_case = {str(r.get("case_id")): r for r in sheet_rows(wb, "NORMALIZATION")}

    for r in sheet_rows(wb, "DEVICE_CAPACITY"):
        cid = str(r.get("case_id"))
        case = case_by_id.get(cid)
        if not case:
            continue
        kind = str(r.get("kind"))
        item = find_device(case["system"], kind, int(r.get("component_index") or -1), int(r.get("position") or -1))
        if item is not None:
            field = str(r.get("capacity_field") or "")
            base_mw = float(r.get("base_mw") or 0.0)
            if field:
                item[field] = base_mw
            # The solvers/modules do not all use the same capacity field. Keep
            # companion fields in sync so editing one Excel base_mw row has an
            # actual effect in time-series and resilience modules.
            if kind == "AC_PV_SYSTEM":
                item["pmax_mw"] = base_mw
                item["p_mw"] = base_mw
                item["sn_mva"] = max(float(item.get("sn_mva") or 0.0), base_mw)
            elif kind in ("AC_WIND", "AC_RENEWABLE"):
                item["p_rated_mw"] = base_mw
                item["p_mw"] = base_mw
            elif kind in ("AC_STATIC_WIND", "AC_STATIC_PV"):
                item["p_rated_mw"] = base_mw
                item["pmax_mw"] = base_mw
                item["p_mw"] = base_mw
            elif kind == "DC_PV_ARRAY":
                item["p_set_mw"] = base_mw
            elif kind in ("DC_STATIC_WIND", "DC_STATIC_PV"):
                item["p_set_mw"] = base_mw
                item["pmax_mw"] = base_mw
            if "e_rated_mwh" in item:
                item["e_rated_mwh"] = float(r.get("base_mwh") or item.get("e_rated_mwh") or 0.0)

    for case in norm["cases"]:
        cid = case["case_id"]
        profiles = list(catalog.get(cid, {}).values())
        ts = case["standard_time_series"] or {}
        if profiles:
            ts["profiles"] = profiles
            ts["num_steps"] = len(profiles[0].get("values") or [])
        if cid in binding_by_case:
            ts["binding"] = binding_by_case[cid]
        if cid in norm_by_case:
            n = norm_by_case[cid]
            ts["step_duration_hr"] = float(n.get("step_duration_hr") or ts.get("step_duration_hr") or 1.0)
            ts["normalization"] = {
                "base_load_mw": float(n.get("base_load_mw") or 0.0),
                "base_pv_mw": float(n.get("base_pv_mw") or 0.0),
                "base_wind_mw": float(n.get("base_wind_mw") or 0.0),
                "base_other_renewable_mw": float(n.get("base_other_renewable_mw") or 0.0),
                "base_renewable_mw": float(n.get("base_renewable_mw") or 0.0),
                "profile_semantics": n.get("profile_semantics") or "dimensionless_multiplier",
            }
        case["standard_time_series"] = ts

    # Reconstruct original v2-style generated scenario cases because those are
    # what the existing GUI import/export flow expects.
    out_cases = []
    for case in norm["cases"]:
        system = dict(case["system"])
        system["_generated_scenario"] = case["generated_scenario"]
        system["_time_series"] = case["standard_time_series"]
        out_cases.append(system)

    out = dict(bundle)
    out["cases"] = out_cases
    out["case_count"] = len(out_cases)
    if "format" not in out:
        out["format"] = "generated_scenario_case_bundle_v2"
    if "schema_version" not in out:
        out["schema_version"] = 2
    json_path.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=["json-to-xlsx", "xlsx-to-json"])
    parser.add_argument("input")
    parser.add_argument("output")
    args = parser.parse_args()
    inp = Path(args.input)
    out = Path(args.output)
    if args.mode == "json-to-xlsx":
        json_to_xlsx(inp, out)
    else:
        xlsx_to_json(inp, out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
